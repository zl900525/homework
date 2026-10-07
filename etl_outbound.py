# /// script
# dependencies = [
#     "pandas",
#     "openpyxl",
# ]
# ///

import os
import re
import pandas as pd
import openpyxl

def clean_name(text: str) -> str:
    """清理洲別與國家名稱中的中英文混雜、換行符號與多餘空格"""
    if not text:
        return ""
    first_line = str(text).split('\n')[0].strip()
    zh_parts = re.findall(r'[一-龥]+', first_line)
    if zh_parts:
        return "".join(zh_parts)
    return first_line

def to_int(val):
    """數值轉 int；空白或非數值回傳 None（不要默默當成 0）"""
    if val is None or val == "":
        return None
    try:
        return int(val)
    except (ValueError, TypeError):
        return None

def read_records(input_file: str) -> list[dict]:
    wb = openpyxl.load_workbook(input_file, data_only=True)
    ws = wb.worksheets[0]

    region_map = {}
    current_region = "其他"

    # 讀取第 2 列：區域 / 洲別（處理跨欄合併與名稱清洗）
    for c in range(2, ws.max_column + 1):
        r_val = ws.cell(row=2, column=c).value
        if r_val is not None:
            cleaned_r = clean_name(str(r_val))
            if cleaned_r:
                current_region = cleaned_r
        region_map[c] = current_region

    # 讀取第 3 列：細分國家 / 地區名稱
    country_map = {}
    for c in range(2, ws.max_column + 1):
        c_val = ws.cell(row=3, column=c).value
        if c_val is not None:
            country_map[c] = clean_name(str(c_val))
        else:
            c_val_r2 = ws.cell(row=2, column=c).value
            if c_val_r2:
                country_map[c] = clean_name(str(c_val_r2))

    # 「X洲其他地區」欄在原始檔是公式（= 該洲合計 - 列出國家加總），
    # 檔案未存公式快取值，data_only=True 會讀到 None，因此改由合計欄自行回推
    others_formula = {}
    for c, name in country_map.items():
        if "其他地區" not in name:
            continue
        region = region_map[c]
        same_region = [k for k in country_map if region_map[k] == region]
        total_col = next(k for k in same_region if "合計" in country_map[k])
        listed_cols = [k for k in same_region if k not in (c, total_col)]
        others_formula[c] = (total_col, listed_cols)

    records = []

    # 讀取歷年數據 (Row 4 起至表尾)
    for r in range(4, ws.max_row + 1):
        year_cell = ws.cell(row=r, column=1).value
        if not year_cell or "註" in str(year_cell) or "資料來源" in str(year_cell):
            continue

        year_str = str(year_cell).strip()
        match = re.search(r'(\d+)年', year_str)
        if not match:
            continue
        year_roc = int(match.group(1))

        row_vals = {c: to_int(ws.cell(row=r, column=c).value) for c in country_map}
        for c, (total_col, listed_cols) in others_formula.items():
            if row_vals[c] is None and row_vals[total_col] is not None:
                row_vals[c] = row_vals[total_col] - sum(row_vals[k] or 0 for k in listed_cols)

        for c, country_raw in country_map.items():
            if not country_raw:
                continue
            records.append({
                "year": year_roc,
                "region": region_map[c],
                "country": country_raw,
                "count": row_vals[c],
            })

    return records

def etl_outbound():
    input_file = "表2-2-歷年中華民國國民出國目的地人次統計.xlsx"
    output_csv = "work/outbound_stat.csv"

    os.makedirs(os.path.dirname(output_csv), exist_ok=True)

    out_df = pd.DataFrame(read_records(input_file))

    # 排除「合計」、「小計」、「總計」等重複計算列
    out_df = out_df[~out_df["country"].str.contains("合計|小計|總計", na=False)]

    # 同一年份、同一區域與國家進行分組加總
    out_df = out_df.groupby(["year", "region", "country"], as_index=False)["count"].sum(min_count=1)
    out_df["count"] = out_df["count"].astype("Int64")

    # 存至 CSV，採用 UTF-8 with BOM 格式
    out_df.to_csv(output_csv, index=False, encoding="utf-8-sig")
    print(f"ETL 成功，已產出：{output_csv}")

if __name__ == "__main__":
    etl_outbound()
